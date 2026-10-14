#!/usr/bin/env python3
"""Build an offline editable PowerPoint and independent preview/report artifacts."""

from __future__ import annotations

import argparse
import json
import shutil
import sys
import traceback
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import yaml
from pptx.enum.text import PP_ALIGN

from editable_scene import (DEFAULT_THEME, Model, NativeSlide, ShapeCounts, crystal_cluster,
                            hex_rgb, molecule_from_file, molecule_from_smiles, new_presentation,
                            normalized_xy)

OUTPUT_FILES = ("editable_mechanism.pptx", "preview.png", "structure_report.md", "build_log.txt", "shape_manifest.yaml")


def safe_output_dir(path: Path, force: bool) -> Path:
    path = path.resolve()
    if path.exists():
        present = [path / name for name in OUTPUT_FILES if (path / name).exists()]
        if present and not force:
            raise FileExistsError(f"output files already exist in {path}; use --force")
        if not path.is_dir():
            raise ValueError(f"output path is not a directory: {path}")
    else:
        path.mkdir(parents=True)
    return path


def resolve_input(value: str, base: Path) -> Path:
    path = Path(value)
    if not path.is_absolute():
        path = base / path
    path = path.resolve()
    if not path.is_file():
        raise FileNotFoundError(path)
    return path


def molecule_spec(spec: dict[str, Any], base: Path) -> Model:
    label = str(spec.get("label", spec.get("id", "Molecule")))
    if "smiles" in spec:
        return molecule_from_smiles(str(spec["smiles"]), label)
    if "input" in spec:
        return molecule_from_file(resolve_input(str(spec["input"]), base), label)
    raise ValueError(f"species {label!r} needs smiles or input")


def add_title(native: NativeSlide, title: str):
    native.text(title, 0.45, 0.18, 12.45, 0.62, size=24, bold=True, align=PP_ALIGN.LEFT, name="title")


def draw_mechanism_slide(prs, slide_spec: dict[str, Any], base: Path, theme: dict[str, str],
                         counts: ShapeCounts, models: list[Model]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = hex_rgb(theme["background"])
    native = NativeSlide(slide, counts, theme)
    add_title(native, str(slide_spec.get("title", "Editable mechanism")))
    if "catalyst" in slide_spec:
        native.catalyst(dict(slide_spec["catalyst"]))
    boxes: dict[str, tuple[float, float, float, float]] = {}
    species = slide_spec.get("species", [])
    if len(species) > 12:
        raise ValueError("mechanism slide has more than 12 species")
    for spec in species:
        model = molecule_spec(spec, base); models.append(model)
        x, y, w, h = (float(spec[k]) for k in ("x", "y", "width", "height"))
        boxes[str(spec["id"])] = (x, y, w, h)
        native.model(model, x, y, w, h, prefix=str(spec["id"]))
        native.text(str(spec.get("label", spec["id"])), x, y + h + 0.05, w, 0.38, size=12, bold=True, name=f"species_{spec['id']}")
    for i, arrow in enumerate(slide_spec.get("arrows", [])):
        if arrow["from"] not in boxes or arrow["to"] not in boxes:
            raise ValueError(f"arrow endpoint not found: {arrow}")
        left, right = boxes[arrow["from"]], boxes[arrow["to"]]
        x1, y1 = left[0] + left[2] + 0.08, left[1] + left[3] / 2
        x2, y2 = right[0] - 0.08, right[1] + right[3] / 2
        native.line(x1, y1, x2, y2, name=f"Arrow_step_{i}", color=theme["accent"], width=2.1, arrow=True)
        if arrow.get("label"):
            label_width = max(1.45, abs(x2 - x1))
            midpoint = (x1 + x2) / 2
            native.text(str(arrow["label"]), midpoint - label_width / 2, min(y1, y2) - 0.68,
                        label_width, 0.48, size=9, color=theme["accent"], name=f"arrow_{i}")
    for i, item in enumerate(slide_spec.get("texts", [])):
        native.text(str(item["text"]), float(item["x"]), float(item["y"]), float(item["width"]),
                    float(item["height"]), size=float(item.get("size", 12)), name=f"note_{i}")
    if "crystal" in slide_spec:
        spec = dict(slide_spec["crystal"])
        model = crystal_cluster(resolve_input(str(spec["input"]), base), str(spec.get("center_element", "Zr")), float(spec.get("radius", 3.2)))
        models.append(model)
        x, y, w, h = (float(spec[k]) for k in ("x", "y", "width", "height"))
        native.model(model, x, y, w, h, prefix="crystal")
        native.text(str(spec.get("label", model.label)), x, y + h, w, 0.42, size=12, bold=True, name="crystal_label")


def preview_models(models: list[Model], output: Path, title: str):
    count = max(1, len(models)); cols = min(3, count); rows = (count + cols - 1) // cols
    fig, axes = plt.subplots(rows, cols, figsize=(4.4 * cols, 3.4 * rows), squeeze=False)
    for ax in axes.ravel(): ax.axis("off")
    for ax, model in zip(axes.ravel(), models):
        xy = normalized_xy(model.atoms)
        for bond in model.bonds:
            ax.plot([xy[bond.a, 0], xy[bond.b, 0]], [xy[bond.a, 1], xy[bond.b, 1]], color="#4B5563", lw=1.3, zorder=1)
        for element in sorted({a.element for a in model.atoms}):
            idx = [i for i, a in enumerate(model.atoms) if a.element == element]
            colors = {"H":"#F3F4F6","C":"#374151","O":"#DC2626","V":"#7C3AED","Zr":"#14B8A6"}
            ax.scatter(xy[idx, 0], xy[idx, 1], s=100, c=colors.get(element, "#9CA3AF"), edgecolors="#111827", zorder=2, label=element)
        ax.set_title(model.label); ax.set_aspect("equal"); ax.margins(0.18)
    fig.suptitle(title, fontsize=15); fig.tight_layout(); fig.savefig(output, dpi=160, bbox_inches="tight"); plt.close(fig)


def report(models: list[Model], counts: ShapeCounts, source: str) -> str:
    lines = ["# Structure report", "", f"- Generated: {datetime.now(timezone.utc).isoformat()}",
             f"- Source: `{source}`", f"- Models: {len(models)}", f"- Native atoms: {counts.atoms}",
             f"- Native bond segments: {counts.bonds}", f"- Native arrows: {counts.arrows}",
             f"- Independent text boxes: {counts.texts}", f"- Catalyst/support shapes: {counts.catalysts}",
             "- Embedded pictures: 0", "", "## Models", ""]
    for i, model in enumerate(models, 1):
        composition: dict[str, int] = {}
        for atom in model.atoms: composition[atom.element] = composition.get(atom.element, 0) + 1
        lines += [f"### {i}. {model.label}", "", f"- Atoms: {len(model.atoms)}", f"- Bonds: {len(model.bonds)}",
                  "- Composition: " + ", ".join(f"{k}: {v}" for k, v in sorted(composition.items()))]
        lines += [f"- {note}" for note in model.notes] + [""]
    lines += ["## Interpretation limits", "", "Coordinates are projected into two dimensions for communication. "
              "Inferred bonds and supplied mechanism arrows require scientific review; this file is not a refinement or kinetic model.", ""]
    return "\n".join(lines)


def build(args) -> dict[str, Any]:
    out = safe_output_dir(Path(args.output_dir), args.force)
    prs = new_presentation(); counts = ShapeCounts(); models: list[Model] = []
    theme = dict(DEFAULT_THEME)
    source = getattr(args, "smiles", None) or getattr(args, "input", None) or "direct input"
    if args.command == "mechanism":
        input_path = Path(args.input).resolve(); data = yaml.safe_load(input_path.read_text())
        if not isinstance(data, dict) or data.get("schema_version") != 1:
            raise ValueError("mechanism YAML must be a mapping with schema_version: 1")
        theme.update({k: str(v) for k, v in data.get("theme", {}).items() if k in theme})
        slides = data.get("slides", [])
        if not slides or len(slides) > 12: raise ValueError("mechanism YAML requires 1-12 slides")
        for slide_spec in slides: draw_mechanism_slide(prs, slide_spec, input_path.parent, theme, counts, models)
        title = str(data.get("title", input_path.stem))
    else:
        slide = prs.slides.add_slide(prs.slide_layouts[6]); native = NativeSlide(slide, counts, theme)
        title = args.title or (Path(args.input).stem if args.input else "Editable molecule")
        add_title(native, title)
        if args.command == "molecule":
            model = molecule_from_smiles(args.smiles, title) if args.smiles else molecule_from_file(Path(args.input).resolve(), title)
        else:
            model = crystal_cluster(Path(args.input).resolve(), args.center_element, args.radius)
        models.append(model); native.model(model, 1.0, 1.0, 11.3, 5.7, prefix=args.command)
        native.text(model.label, 3.0, 6.7, 7.3, 0.38, size=12, bold=True, name="model_label")
    prs.save(out / "editable_mechanism.pptx")
    preview_models(models, out / "preview.png", title)
    (out / "structure_report.md").write_text(report(models, counts, str(source)))
    manifest = {"contract_version": 1, "expected": counts.as_dict(), "pictures": 0, "groups": 0,
                "slides": len(prs.slides), "models": len(models)}
    (out / "shape_manifest.yaml").write_text(yaml.safe_dump(manifest, sort_keys=False))
    return {"output_dir": str(out), "manifest": manifest, "models": len(models)}


def parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description=__doc__)
    sub = p.add_subparsers(dest="command", required=True)
    common = argparse.ArgumentParser(add_help=False)
    common.add_argument("--output-dir", default="outputs"); common.add_argument("--force", action="store_true")
    common.add_argument("--title")
    m = sub.add_parser("molecule", parents=[common]); group = m.add_mutually_exclusive_group(required=True)
    group.add_argument("--smiles"); group.add_argument("--input")
    c = sub.add_parser("crystal", parents=[common]); c.add_argument("--input", required=True)
    c.add_argument("--center-element", default="Zr"); c.add_argument("--radius", type=float, default=3.2)
    mech = sub.add_parser("mechanism", parents=[common]); mech.add_argument("--input", required=True)
    return p


def main() -> int:
    args = parser().parse_args(); out = Path(args.output_dir).resolve(); log_path = out / "build_log.txt"
    try:
        result = build(args)
        log = ["editable-catalysis-mechanism-pptx build", f"status: PASS", f"command: {args.command}",
               f"source: {getattr(args, 'smiles', None) or getattr(args, 'input', None)}", f"output_dir: {result['output_dir']}",
               "network_api_calls: 0", "embedded_pictures: 0", json.dumps(result["manifest"], sort_keys=True)]
        log_path.write_text("\n".join(log) + "\n")
        print(json.dumps(result, indent=2)); return 0
    except Exception as exc:
        out.mkdir(parents=True, exist_ok=True)
        log_path.write_text(f"editable-catalysis-mechanism-pptx build\nstatus: FAIL\nerror: {exc}\n{traceback.format_exc()}")
        print(f"ERROR: {exc}", file=sys.stderr); return 1


if __name__ == "__main__":
    raise SystemExit(main())
