#!/usr/bin/env python3
"""Validate that a PPTX uses expected independent native shapes."""

from __future__ import annotations

import argparse
import json
import zipfile
from collections import Counter
from pathlib import Path

import yaml
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE


PREFIXES = {"Atom_": "atoms", "Bond_": "bonds", "Arrow_": "arrows", "Text_": "texts", "Catalyst_": "catalysts"}


def validate(path: Path, manifest_path: Path | None = None) -> dict:
    prs = Presentation(path)
    counts = Counter({value: 0 for value in PREFIXES.values()})
    pictures = groups = unnamed = 0
    type_errors: list[str] = []
    duplicate_names: set[str] = set()
    for slide in prs.slides:
        slide_names: list[str] = []
        for shape in slide.shapes:
            slide_names.append(shape.name)
            pictures += int(shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
            groups += int(shape.shape_type == MSO_SHAPE_TYPE.GROUP)
            matched = False
            for prefix, category in PREFIXES.items():
                if shape.name.startswith(prefix):
                    counts[category] += 1; matched = True
                    if prefix == "Atom_" and (shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE or shape.auto_shape_type != MSO_AUTO_SHAPE_TYPE.OVAL):
                        type_errors.append(f"{shape.name} is not a native oval")
                    elif prefix in {"Bond_", "Arrow_"} and shape.shape_type != MSO_SHAPE_TYPE.LINE:
                        type_errors.append(f"{shape.name} is not a native connector")
                    elif prefix == "Text_" and not shape.has_text_frame:
                        type_errors.append(f"{shape.name} is not an independent text shape")
                    if prefix == "Arrow_":
                        xml = shape.element.xml
                        if 'type="triangle"' not in xml:
                            type_errors.append(f"{shape.name} has no native triangle arrowhead")
                    break
            if not matched: unnamed += 1
        duplicate_names.update(name for name, n in Counter(slide_names).items() if n > 1)
    duplicates = sorted(duplicate_names)
    with zipfile.ZipFile(path) as archive:
        media = [name for name in archive.namelist() if name.startswith("ppt/media/")]
        slide_xml = b"".join(archive.read(name) for name in archive.namelist()
                             if name.startswith("ppt/slides/slide") and name.endswith(".xml"))
        picture_xml_count = slide_xml.count(b"<p:pic>")
        group_xml_count = slide_xml.count(b"<p:grpSp>")
    errors: list[str] = []
    if pictures or picture_xml_count or media:
        errors.append(f"pictures forbidden: object={pictures}, xml={picture_xml_count}, media={len(media)}")
    if groups or group_xml_count:
        errors.append(f"group shapes forbidden: object={groups}, xml={group_xml_count}")
    if duplicates: errors.append(f"duplicate shape names: {duplicates[:8]}")
    errors.extend(type_errors)
    if counts["atoms"] < 1: errors.append("no Atom_* native shapes")
    if counts["bonds"] < 1: errors.append("no Bond_* native connectors")
    if counts["texts"] < 1: errors.append("no Text_* independent text boxes")
    expected = None
    if manifest_path:
        expected = yaml.safe_load(manifest_path.read_text())
        for key, value in expected.get("expected", {}).items():
            if counts[key] != value: errors.append(f"{key}: expected {value}, found {counts[key]}")
        if len(prs.slides) != expected.get("slides"): errors.append("slide count differs from manifest")
    result = {"status": "PASS" if not errors else "FAIL", "file": str(path), "slides": len(prs.slides),
              "native_shapes": dict(counts), "pictures": pictures, "groups": groups,
              "media_files": media, "unclassified_native_shapes": unnamed, "errors": errors}
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", type=Path); parser.add_argument("--expect-manifest", type=Path)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()
    result = validate(args.pptx, args.expect_manifest)
    print(json.dumps(result, indent=2) if args.json else yaml.safe_dump(result, sort_keys=False))
    return 0 if result["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
