"""Offline scene construction and native PowerPoint rendering utilities."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
import math

import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ELEMENT_COLORS = {
    "H": "#F3F4F6", "C": "#374151", "N": "#2563EB", "O": "#DC2626",
    "F": "#16A34A", "P": "#F97316", "S": "#EAB308", "Cl": "#22C55E",
    "V": "#7C3AED", "Zr": "#14B8A6",
}
ELEMENT_RADII = {"H": 0.15, "C": 0.22, "N": 0.22, "O": 0.22, "V": 0.30, "Zr": 0.34}


def hex_rgb(value: str) -> RGBColor:
    value = value.strip().lstrip("#")
    if len(value) != 6:
        raise ValueError(f"invalid RGB color: {value!r}")
    return RGBColor.from_string(value.upper())


@dataclass
class Atom:
    element: str
    x: float
    y: float
    z: float = 0.0
    source_index: int = 0
    image: tuple[int, int, int] = (0, 0, 0)


@dataclass
class Bond:
    a: int
    b: int
    order: int = 1


@dataclass
class Model:
    label: str
    atoms: list[Atom]
    bonds: list[Bond]
    notes: list[str] = field(default_factory=list)


@dataclass
class ShapeCounts:
    atoms: int = 0
    bonds: int = 0
    arrows: int = 0
    texts: int = 0
    catalysts: int = 0

    def as_dict(self) -> dict[str, int]:
        return {k: int(v) for k, v in vars(self).items()}


def normalized_xy(atoms: list[Atom]) -> np.ndarray:
    if not atoms:
        return np.empty((0, 2))
    points = np.asarray([[a.x, a.y, a.z] for a in atoms], dtype=float)
    points -= points.mean(axis=0)
    if len(points) >= 2:
        _, _, vh = np.linalg.svd(points, full_matrices=False)
        xy = points @ vh[:2].T
    else:
        xy = points[:, :2]
    span = np.ptp(xy, axis=0)
    span[span < 1e-8] = 1.0
    return xy / max(float(span.max()), 1e-8)


def molecule_from_smiles(smiles: str, label: str = "Molecule", *, max_atoms: int = 200) -> Model:
    from rdkit import Chem
    from rdkit.Chem import AllChem

    mol = Chem.MolFromSmiles(smiles)
    if mol is None:
        raise ValueError(f"RDKit could not parse SMILES: {smiles!r}")
    mol = Chem.AddHs(mol)
    if mol.GetNumAtoms() > max_atoms:
        raise ValueError(f"molecule has {mol.GetNumAtoms()} atoms; limit is {max_atoms}")
    params = AllChem.ETKDGv3()
    params.randomSeed = 0xC0D3
    status = AllChem.EmbedMolecule(mol, params)
    if status != 0:
        AllChem.Compute2DCoords(mol)
    elif AllChem.UFFHasAllMoleculeParams(mol):
        AllChem.UFFOptimizeMolecule(mol, maxIters=250)
    conf = mol.GetConformer()
    atoms = [Atom(a.GetSymbol(), conf.GetAtomPosition(i).x, conf.GetAtomPosition(i).y,
                  conf.GetAtomPosition(i).z, i) for i, a in enumerate(mol.GetAtoms())]
    bonds = [Bond(b.GetBeginAtomIdx(), b.GetEndAtomIdx(), max(1, int(round(b.GetBondTypeAsDouble()))))
             for b in mol.GetBonds()]
    return Model(label, atoms, bonds, [f"Input SMILES: `{smiles}`", "Hydrogens added explicitly by RDKit."])


def molecule_from_file(path: Path, label: str | None = None, *, max_atoms: int = 200) -> Model:
    from rdkit import Chem
    from rdkit.Chem import AllChem, rdDetermineBonds

    suffix = path.suffix.lower()
    notes = [f"Input file: `{path.name}`"]
    if suffix in {".smi", ".smiles"}:
        first = next((line.strip() for line in path.read_text().splitlines() if line.strip()), "")
        if not first:
            raise ValueError("SMILES file is empty")
        fields = first.split()
        return molecule_from_smiles(fields[0], label or (fields[1] if len(fields) > 1 else path.stem), max_atoms=max_atoms)
    if suffix in {".sdf", ".mol"}:
        mol = Chem.MolFromMolFile(str(path), removeHs=False, sanitize=True)
    elif suffix == ".xyz":
        mol = Chem.MolFromXYZFile(str(path))
        if mol is not None:
            rdDetermineBonds.DetermineBonds(mol, charge=0)
            notes.append("Bonds inferred by RDKit DetermineBonds with total charge 0.")
    else:
        raise ValueError(f"unsupported molecule format: {suffix}")
    if mol is None:
        raise ValueError(f"RDKit could not parse {path}")
    if mol.GetNumAtoms() > max_atoms:
        raise ValueError(f"molecule has {mol.GetNumAtoms()} atoms; limit is {max_atoms}")
    if mol.GetNumConformers() == 0:
        mol = Chem.AddHs(mol)
        params = AllChem.ETKDGv3(); params.randomSeed = 0xC0D3
        if AllChem.EmbedMolecule(mol, params) != 0:
            AllChem.Compute2DCoords(mol)
    conf = mol.GetConformer()
    atoms = [Atom(a.GetSymbol(), conf.GetAtomPosition(i).x, conf.GetAtomPosition(i).y,
                  conf.GetAtomPosition(i).z, i) for i, a in enumerate(mol.GetAtoms())]
    bonds = [Bond(b.GetBeginAtomIdx(), b.GetEndAtomIdx(), max(1, int(round(b.GetBondTypeAsDouble()))))
             for b in mol.GetBonds()]
    return Model(label or path.stem, atoms, bonds, notes)


def crystal_cluster(path: Path, center_element: str = "Zr", radius: float = 3.2,
                    *, max_atoms: int = 120) -> Model:
    from pymatgen.core import Structure

    structure = Structure.from_file(path)
    centers = [i for i, site in enumerate(structure) if site.specie.symbol == center_element]
    if not centers:
        raise ValueError(f"no {center_element} site found in {path.name}")
    center_index = centers[0]
    center = structure[center_index]
    neighbors = structure.get_sites_in_sphere(center.coords, radius, include_index=True, include_image=True)
    if len(neighbors) > max_atoms:
        raise ValueError(f"cluster has {len(neighbors)} atoms; limit is {max_atoms}")
    neighbors.sort(key=lambda n: (float(n.nn_distance), int(n.index), tuple(n.image)))
    atoms: list[Atom] = []
    for n in neighbors:
        delta = np.asarray(n.coords) - np.asarray(center.coords)
        atoms.append(Atom(n.specie.symbol, *map(float, delta), int(n.index), tuple(map(int, n.image))))
    bonds: list[Bond] = []
    covalent = {"O": 0.66, "Zr": 1.48, "V": 1.22, "H": 0.31, "C": 0.76}
    for i in range(len(atoms)):
        for j in range(i + 1, len(atoms)):
            distance = math.dist((atoms[i].x, atoms[i].y, atoms[i].z), (atoms[j].x, atoms[j].y, atoms[j].z))
            cutoff = 1.18 * (covalent.get(atoms[i].element, 0.8) + covalent.get(atoms[j].element, 0.8))
            if 0.3 < distance <= cutoff:
                bonds.append(Bond(i, j, 1))
    notes = [f"Input structure: `{path.name}`", f"Formula: {structure.composition.reduced_formula}",
             f"Space group: {structure.get_space_group_info()[0]}",
             f"Cluster center: first {center_element} site (index {center_index})",
             f"Periodic radius cutoff: {radius:.2f} Å", "Bonds inferred from 1.18 × covalent-radius sums."]
    return Model(f"{center_element}-centered cluster", atoms, bonds, notes)


class NativeSlide:
    def __init__(self, slide, counts: ShapeCounts, theme: dict[str, str]):
        self.slide = slide
        self.counts = counts
        self.theme = theme

    def text(self, text: str, x: float, y: float, w: float, h: float, *,
             size: float = 14, bold: bool = False, color: str | None = None,
             align=PP_ALIGN.CENTER, name: str = "Text"):
        box = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        box.name = f"Text_{name}_{self.counts.texts + 1}"
        frame = box.text_frame; frame.clear(); frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = frame.paragraphs[0]; p.alignment = align
        run = p.add_run(); run.text = str(text); run.font.size = Pt(size); run.font.bold = bold
        run.font.name = "Aptos"; run.font.color.rgb = hex_rgb(color or self.theme["foreground"])
        self.counts.texts += 1
        return box

    def line(self, x1: float, y1: float, x2: float, y2: float, *, name: str,
             color: str = "#4B5563", width: float = 1.5, arrow: bool = False):
        shape = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        shape.name = name
        shape.line.color.rgb = hex_rgb(color); shape.line.width = Pt(width)
        if arrow:
            from pptx.oxml.xmlchemy import OxmlElement
            head = OxmlElement("a:headEnd"); head.set("type", "none")
            end = OxmlElement("a:tailEnd"); end.set("type", "triangle")
            shape.line._get_or_add_ln().append(head); shape.line._get_or_add_ln().append(end)
            self.counts.arrows += 1
        else:
            self.counts.bonds += 1
        return shape

    def atom(self, atom: Atom, x: float, y: float, diameter: float, *, index: int, prefix: str = "model"):
        d = max(0.13, diameter * (ELEMENT_RADII.get(atom.element, 0.23) / 0.23))
        shape = self.slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x - d / 2), Inches(y - d / 2), Inches(d), Inches(d))
        shape.name = f"Atom_{prefix}_{atom.element}_{index}"
        shape.fill.solid(); shape.fill.fore_color.rgb = hex_rgb(ELEMENT_COLORS.get(atom.element, "#9CA3AF"))
        shape.line.color.rgb = hex_rgb("#111827"); shape.line.width = Pt(0.8)
        self.counts.atoms += 1
        self.text(atom.element, x - d * 0.55, y - d * 0.42, d * 1.1, d * 0.84,
                  size=max(7, d * 19), bold=True, color="#FFFFFF" if atom.element == "C" else "#111827",
                  name=f"atom_{prefix}_{atom.element}_{index}")
        return shape

    def model(self, model: Model, x: float, y: float, w: float, h: float, *, prefix: str):
        xy = normalized_xy(model.atoms)
        if len(xy) == 0:
            return
        margin = 0.3
        scale = min(max(w - 2 * margin, 0.1), max(h - 2 * margin, 0.1))
        positions = [(x + w / 2 + float(p[0]) * scale * 0.82,
                      y + h / 2 - float(p[1]) * scale * 0.82) for p in xy]
        diameter = max(0.16, min(0.40, min(w, h) / max(6.0, math.sqrt(len(model.atoms)) * 2.3)))
        for b_index, bond in enumerate(model.bonds):
            x1, y1 = positions[bond.a]; x2, y2 = positions[bond.b]
            dx, dy = x2 - x1, y2 - y1; length = math.hypot(dx, dy) or 1.0
            nx, ny = -dy / length, dx / length
            for stroke in range(min(bond.order, 3)):
                offset = (stroke - (min(bond.order, 3) - 1) / 2) * 0.055
                self.line(x1 + nx * offset, y1 + ny * offset, x2 + nx * offset, y2 + ny * offset,
                          name=f"Bond_{prefix}_{b_index}_{stroke}", width=1.4)
        for i, atom in sorted(enumerate(model.atoms), key=lambda item: item[1].z):
            self.atom(atom, *positions[i], diameter, index=i, prefix=prefix)

    def catalyst(self, spec: dict[str, Any]):
        x, y = float(spec.get("x", 0.8)), float(spec.get("y", 5.7))
        w, h = float(spec.get("width", 11.7)), float(spec.get("height", 0.65))
        support = self.slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        support.name = f"Catalyst_support_{self.counts.catalysts + 1}"
        support.fill.solid(); support.fill.fore_color.rgb = hex_rgb(self.theme["catalyst"])
        support.line.color.rgb = hex_rgb("#6B7280"); support.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        self.counts.catalysts += 1
        for i, site in enumerate(spec.get("sites", [])):
            sx = float(site.get("x", x + (i + 1) * w / (len(spec["sites"]) + 1)))
            atom = Atom(str(site["element"]), 0, 0)
            self.atom(atom, sx, y - 0.08, 0.36, index=1000 + i, prefix="catalyst")
        self.text(spec.get("label", "Catalyst"), x, y + 0.08, w, h * 0.7, size=12, bold=True, name="catalyst")


def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


DEFAULT_THEME = {"background": "#FFFFFF", "foreground": "#1F2937", "accent": "#C2410C", "catalyst": "#D1D5DB"}
